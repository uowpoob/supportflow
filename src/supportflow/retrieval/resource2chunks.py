"""policies/data 的 Markdown、PDF、DOCX 切分入库；RetrievedChunk 保持不变。

解析仅访问本地文件。--replace 在向量生成和校验成功后清空知识集合，
再写入本次所有资料；普通运行仅更新同 ID 的片段。
"""

import argparse
import asyncio
import hashlib
import json
import re
import warnings
from pathlib import Path

from langchain_core.documents import Document

RETRIEVAL_DIR = Path(__file__).resolve().parent
DOCUMENT_DIRS = (RETRIEVAL_DIR / "policies", RETRIEVAL_DIR / "data")
SUPPORTED_SUFFIXES = {".md", ".pdf", ".docx"}
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 120
HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)\s*$")
ARTICLE = re.compile(r"^(第[一二三四五六七八九十百千万零〇两\d]+条)(?:\s+(.*))?$")
CHAPTER = re.compile(r"^第[一二三四五六七八九十百千万零〇两\d]+[章节]\s*.*$")


def _read_markdown(path: Path) -> str:
    text = path.read_text(encoding="utf-8-sig").strip()
    if text.startswith("---\n"):
        match = re.match(r"\A---\n.*?\n---(?:\n|$)", text, re.DOTALL)
        if not match:
            raise ValueError(f"Markdown 元信息未闭合：{path}")
        text = text[match.end():].strip()
    if not text.startswith("# "):
        raise ValueError(f"文档为空或缺少一级标题：{path}")
    return text


def _read_docx(path: Path) -> str:
    from docx import Document as WordDocument
    from docx.table import Table

    document = WordDocument(path)
    title_block = next((b for b in document.iter_inner_content() if not isinstance(b, Table)
                        and b.style.name == "Title" and b.text.strip()), None)
    title = title_block.text.strip() if title_block is not None else path.stem

    def render(container):
        for block in container.iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    yield " | ".join("\n".join(render(cell)) for cell in row.cells)
                continue
            text = block.text.strip()
            if not text or block._element is getattr(title_block, "_element", None):
                continue
            heading = re.fullmatch(r"Heading ([1-9])", block.style.name)
            if ARTICLE.fullmatch(text):
                yield f"#### {text}"
            elif heading:
                yield f"{'#' * min(int(heading[1]) + 1, 6)} {text}"
            else:
                yield text

    return f"# {title}\n\n" + "\n\n".join(render(document))


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError(f"PDF 已加密，无法提取：{path}")
    pages, empty_pages = [], []
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            empty_pages.append(number)
            continue
        lines = []
        for line in text.splitlines():
            line = line.strip()
            article = ARTICLE.fullmatch(line)
            if article:
                lines.append(f"#### {article[1]}")
                if article[2]:
                    lines.append(article[2])
            elif CHAPTER.fullmatch(line):
                lines.append(f"### {line}")
            else:
                lines.append(line)
        pages.append(f"## 第{number}页\n\n" + "\n".join(lines))
    if not pages:
        raise ValueError(f"PDF 没有可提取文本，请先进行 OCR：{path}")
    if empty_pages:
        warnings.warn(f"{path} 第 {empty_pages} 页没有文本层；如有正文须先 OCR", stacklevel=2)
    return f"# {path.stem}\n\n" + "\n\n".join(pages)


def load_documents(directories: tuple[Path, ...] = DOCUMENT_DIRS) -> list[Document]:
    """递归读取三种格式；目录、相对路径和扩展名共同构成文档 ID。"""
    readers = {".md": _read_markdown, ".pdf": _read_pdf, ".docx": _read_docx}
    documents, document_ids = [], set()
    for directory in directories:
        directory = Path(directory)
        if not directory.is_dir():
            raise FileNotFoundError(f"知识目录不存在：{directory}")
        paths = sorted(p for p in directory.rglob("*") if p.is_file()
                       and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith("~$"))
        if not paths:
            raise ValueError(f"知识目录中没有 Markdown、PDF 或 DOCX 文件：{directory}")
        for path in paths:
            source = f"{directory.name}/{path.relative_to(directory).as_posix()}"
            if source in document_ids:
                raise ValueError(f"重复的 document_id：{source}")
            document_ids.add(source)
            try:
                text = readers[path.suffix.lower()](path)
            except Exception as exc:
                raise ValueError(f"读取资料失败：{path}：{exc}") from exc
            if not text.partition("\n")[2].strip():
                raise ValueError(f"文档没有正文：{path}")
            documents.append(Document(page_content=text, metadata={
                "document_id": source, "title": text.splitlines()[0][2:].strip(),
                "source": source, "category": directory.name,
            }))
    return documents


def _headings(body: str):
    """忽略代码块中的 #，避免截断示例代码。"""
    offset, fence, previous = 0, None, ""
    for line in body.splitlines(keepends=True):
        marker = re.match(r"^\s{0,3}(`{3,}|~{3,})", line)
        if marker:
            value = marker[1]
            if fence is None:
                fence = value
            elif value[0] == fence[0] and len(value) >= len(fence):
                fence = None
        elif fence is None and (match := HEADING.fullmatch(line.rstrip("\r\n"))):
            # 某些法律原文把“本法第四十二条、第四十三条规定的……”
            # 中的引用单独设成标题；它仍属于上一条正文。
            reference = previous.endswith(("、", "》", "本法", "本条例", "本办法", "和", "及"))
            if not (ARTICLE.fullmatch(match[2]) and reference):
                yield offset, offset + len(line), len(match[1]), match[2]
        if line.strip():
            previous = line.strip()
        offset += len(line)


def _split_text(text: str, size: int, overlap: int) -> list[str]:
    parts, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundaries = [text.rfind(sep, start + size // 2, end) + len(sep)
                          for sep in ("\n\n", "\n", "。", "！", "？", ". ", "; ")]
            valid = [b for b in boundaries if b > start + size // 2]
            if valid:
                end = max(valid)
        content = text[start:end].strip()
        if content:
            parts.append(content)
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return parts


def split_documents(documents: list[Document], *, chunk_size: int = CHUNK_SIZE,
                    chunk_overlap: int = CHUNK_OVERLAP) -> list[Document]:
    """按标题/法律条款切分，超长正文继续切分，标题计入长度上限。"""
    if chunk_size <= 0 or not 0 <= chunk_overlap < chunk_size // 2:
        raise ValueError("chunk_size 必须为正，chunk_overlap 必须小于其一半")
    chunks, chunk_ids = [], set()
    for document in documents:
        body = document.page_content.partition("\n")[2].strip()
        headings = list(_headings(body))
        sections = [("overview", "概述", body[:headings[0][0]] if headings else body, None)]
        hierarchy, page_number = [], None
        for index, (_, start, level, title) in enumerate(headings):
            end = headings[index + 1][0] if index + 1 < len(headings) else len(body)
            while hierarchy and hierarchy[-1][0] >= level:
                hierarchy.pop()
            hierarchy.append((level, title))
            if document.metadata["source"].lower().endswith(".pdf") and level == 2:
                page = re.fullmatch(r"第(\d+)页", title)
                page_number = int(page[1]) if page else None
            section = " / ".join(t for _, t in hierarchy)
            number = re.match(r"\d+(?:\.\d+)*(?=\s|$)", title)
            article = ARTICLE.fullmatch(title)
            section_id = (number[0] if number else article[1] if article else
                          "section-" + hashlib.sha256(section.encode()).hexdigest()[:12])
            if page_number is not None:
                section_id = f"page-{page_number}:{section_id}"
            sections.append((section_id, section, body[start:end], page_number))
        before = len(chunks)
        for section_id, section, content, page_number in sections:
            content = content.strip()
            if not content:
                continue
            prefix = f"# {document.metadata['title']}\n\n## {section}\n\n"
            budget = chunk_size - len(prefix)
            if budget <= 2 * chunk_overlap:
                raise ValueError(f"标题过长，请增大 chunk_size：{document.metadata['source']}")
            parts = _split_text(content, budget, chunk_overlap)
            for index, part in enumerate(parts, 1):
                suffix = f":part-{index}" if len(parts) > 1 else ""
                chunk_id = f"{document.metadata['document_id']}#{section_id}{suffix}"
                if len(chunk_id.encode()) > 512:
                    chunk_id = "sha256:" + hashlib.sha256(chunk_id.encode()).hexdigest()
                if chunk_id in chunk_ids:
                    raise ValueError(f"重复的 chunk_id：{chunk_id}")
                chunk_ids.add(chunk_id)
                source = document.metadata["source"]
                if page_number is not None:
                    source += f"#page={page_number}"
                chunks.append(Document(page_content=prefix + part, metadata={
                    **document.metadata, "chunk_id": chunk_id, "section": section, "source": source,
                }))
        if len(chunks) == before:
            raise ValueError(f"文档没有可切分正文：{document.metadata['source']}")
    return chunks


async def ingest_documents(directories: tuple[Path, ...] = DOCUMENT_DIRS, *,
                           replace: bool = False, chunk_size: int = CHUNK_SIZE,
                           chunk_overlap: int = CHUNK_OVERLAP) -> dict[str, int]:
    from supportflow.core.config import settings
    from supportflow.retrieval.embedding import embed_texts
    from supportflow.retrieval.milvus import clear_chunks, init_milvus, prepare_rows, upsert_chunks

    documents = load_documents(directories)
    chunks = split_documents(documents, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    if not chunks:
        raise ValueError("没有可入库的正文片段")
    if not settings.DASHSCOPE_API_KEY.get_secret_value().strip():
        raise ValueError("请先配置 DASHSCOPE_API_KEY")
    print(f"读取 {len(documents)} 份文档，切分为 {len(chunks)} 个片段", flush=True)
    client = await asyncio.to_thread(init_milvus)
    try:
        print(f"正在向量化：{settings.EMBEDDING_MODEL} / {settings.EMBEDDING_DIM} 维", flush=True)
        vectors = []
        for start in range(0, len(chunks), 100):
            batch = chunks[start:start + 100]
            vectors.extend(await embed_texts([chunk.page_content for chunk in batch]))
            print(f"已向量化 {min(start + 100, len(chunks))}/{len(chunks)} 个片段", flush=True)
        prepare_rows(chunks, vectors)
        deleted = await asyncio.to_thread(clear_chunks, client) if replace else None
        if replace:
            print(f"已清空 {settings.DB_NAME}.{settings.COLLECTION} 的 {deleted} 条旧记录", flush=True)
        written = await asyncio.to_thread(upsert_chunks, client, chunks, vectors)
        stored = 0
        for start in range(0, len(chunks), 100):
            ids = [c.metadata["chunk_id"] for c in chunks[start:start + 100]]
            counts = await asyncio.to_thread(client.query, collection_name=settings.COLLECTION,
                                            filter=f"chunk_id in {json.dumps(ids, ensure_ascii=False)}",
                                            output_fields=["count(*)"], consistency_level="Strong")
            stored += counts[0]["count(*)"]
        if stored != len(chunks):
            raise RuntimeError(f"入库校验失败：期望 {len(chunks)} 条，查询到 {stored} 条")
        if replace:
            counts = await asyncio.to_thread(client.query, collection_name=settings.COLLECTION,
                                            filter="", output_fields=["count(*)"], consistency_level="Strong")
            if counts[0]["count(*)"] != stored:
                raise RuntimeError("重建校验失败：集合仍有本次入库以外的记录")
        summary = {"documents": len(documents), "chunks": len(chunks), "upserted": written, "verified": stored}
        if replace:
            summary["deleted"] = deleted
        return summary
    finally:
        await asyncio.to_thread(client.close)


def main() -> None:
    parser = argparse.ArgumentParser(description="将 policies/data 的 Markdown、PDF、DOCX 切分入库")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="只读取、切分，不调用模型或数据库")
    mode.add_argument("--replace", action="store_true", help="向量化成功后清空知识集合，重新导入所有资料")
    mode.add_argument("--clear-only", action="store_true", help="只清空配置的知识集合，保留集合结构")
    parser.add_argument("--chunk-size", type=int, default=CHUNK_SIZE)
    parser.add_argument("--chunk-overlap", type=int, default=CHUNK_OVERLAP)
    args = parser.parse_args()
    if args.clear_only:
        from supportflow.retrieval.milvus import clear_chunks, init_milvus
        client = init_milvus()
        try:
            print(f"已删除 {clear_chunks(client)} 条知识片段")
        finally:
            client.close()
    elif args.dry_run:
        documents = load_documents()
        chunks = split_documents(documents, chunk_size=args.chunk_size, chunk_overlap=args.chunk_overlap)
        formats = {suffix: sum(d.metadata["source"].lower().endswith(suffix) for d in documents)
                   for suffix in sorted(SUPPORTED_SUFFIXES)}
        print(json.dumps({"documents": len(documents), "chunks": len(chunks), "formats": formats,
                          "categories": {category: sum(c.metadata["category"] == category for c in chunks)
                                         for category in ("policies", "data")}}, ensure_ascii=False))
    else:
        summary = asyncio.run(ingest_documents(replace=args.replace, chunk_size=args.chunk_size,
                                               chunk_overlap=args.chunk_overlap))
        print(f"入库完成：{json.dumps(summary, ensure_ascii=False)}")


if __name__ == "__main__":
    main()
