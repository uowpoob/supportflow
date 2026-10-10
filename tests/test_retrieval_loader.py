"""离线验证来源可追溯、正文不丢失，以及条款 ID 的稳定性。"""

import pytest

from supportflow.retrieval.resource2chunks import load_documents, split_documents


def test_load_and_split_preserve_content_and_context(tmp_path):
    path = tmp_path / "PRD-001.md"
    path.write_text(
        "\ufeff# 示例耳机\n\n适用于第一代。\n\n"
        "## 1.1 连接\n\n先打开蓝牙。\n\n再选择设备。\n\n"
        "## 1.2 维护\n\n### 注意事项\n\n不要进水。\n",
        encoding="utf-8",
    )
    documents = load_documents((tmp_path,))
    chunks = split_documents(documents)

    assert len(documents) == 1
    document_id = f"{tmp_path.name}/PRD-001.md"
    assert [c.metadata["chunk_id"] for c in chunks[:2]] == [
        f"{document_id}#overview", f"{document_id}#1.1",
    ]
    assert "适用于第一代。" in chunks[0].page_content
    assert "先打开蓝牙。\n\n再选择设备。" in chunks[1].page_content
    assert "1.2 维护 / 注意事项\n\n不要进水。" in chunks[2].page_content
    assert all(c.page_content.startswith("# 示例耳机\n") for c in chunks)
    assert all(c.metadata["source"] == f"{tmp_path.name}/PRD-001.md" for c in chunks)
    assert "chunk_id" not in documents[0].metadata

    # 新增较早的条款不会改变既有条款的 ID 或内容。
    path.write_text(path.read_text(encoding="utf-8").replace(
        "## 1.1", "## 1.0 准备\n\n准备手机。\n\n## 1.1",
    ), encoding="utf-8")
    updated = {c.metadata["chunk_id"]: c for c in split_documents(load_documents((tmp_path,)))}
    assert updated[f"{document_id}#1.1"] == chunks[1]


def test_document_without_sections_becomes_overview(tmp_path):
    (tmp_path / "note.md").write_text("# 说明\n\n完整说明。", encoding="utf-8")
    chunks = split_documents(load_documents((tmp_path,)))
    assert len(chunks) == 1
    assert chunks[0].metadata["section"] == "概述"
    assert chunks[0].page_content.endswith("完整说明。")


def test_duplicate_clause_ids_are_rejected(tmp_path):
    (tmp_path / "bad.md").write_text(
        "# 政策\n\n## 1.1 条款甲\n\n甲。\n\n## 1.1 条款乙\n\n乙。",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="重复的 chunk_id"):
        split_documents(load_documents((tmp_path,)))


def test_missing_empty_and_malformed_sources_are_rejected(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_documents((tmp_path / "missing",))
    with pytest.raises(ValueError, match="没有 Markdown"):
        load_documents((tmp_path,))
    (tmp_path / "bad.md").write_text("没有标题", encoding="utf-8")
    with pytest.raises(ValueError, match="缺少一级标题"):
        load_documents((tmp_path,))


def write_pdf(path, pages):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def test_all_formats_same_name_recursive_and_docx_table(tmp_path):
    from docx import Document

    folder = tmp_path / "nested"
    folder.mkdir()
    (folder / "guide.md").write_text(
        '---\nsource_id: "A01"\n---\n\n# Guide\n\n#### 第一条\n\nMarkdown body', encoding="utf-8")
    doc = Document()
    doc.add_paragraph("Guide", style="Title")
    doc.add_heading("Setup", level=1)
    doc.add_paragraph("Word body")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A01"
    table.cell(1, 1).text = "220V"
    doc.add_paragraph("After table")
    doc.save(folder / "guide.docx")
    write_pdf(folder / "guide.PDF", ["PDF page one", "PDF page two"])
    documents = load_documents((tmp_path,))
    chunks = split_documents(documents)
    assert len(documents) == len({d.metadata["document_id"] for d in documents}) == 3
    assert len({c.metadata["chunk_id"] for c in chunks}) == len(chunks)
    word = next(d for d in documents if d.metadata["source"].endswith(".docx"))
    assert word.page_content.index("Word body") < word.page_content.index("A01 | 220V")
    assert word.page_content.index("A01 | 220V") < word.page_content.index("After table")
    md = next(c for c in chunks if c.metadata["source"].endswith(".md"))
    assert md.metadata["section"] == "第一条"
    assert "source_id" not in md.page_content
    pdf = [c for c in chunks if ".PDF#page=" in c.metadata["source"]]
    assert [c.metadata["source"].rsplit("=", 1)[1] for c in pdf] == ["1", "2"]
    assert "PDF page one" in pdf[0].page_content
    assert "PDF page two" in pdf[1].page_content


def test_long_sections_keep_all_text_with_overlap_and_length_bound(tmp_path):
    text = "".join(f"第{i}句是测试内容。" for i in range(500))
    (tmp_path / "long.md").write_text(f"# 长文\n\n## 1 条款\n\n{text}", encoding="utf-8")
    chunks = split_documents(load_documents((tmp_path,)), chunk_size=300, chunk_overlap=30)
    assert len(chunks) > 2
    assert all(len(c.page_content) <= 300 for c in chunks)
    bodies = [c.page_content.split("\n\n", 2)[2] for c in chunks]
    restored = bodies[0] + "".join(body[30:] for body in bodies[1:])
    assert restored == text
    assert chunks == split_documents(load_documents((tmp_path,)), chunk_size=300, chunk_overlap=30)


def test_heading_inside_fenced_code_is_content(tmp_path):
    (tmp_path / "code.md").write_text(
        "# 示例\n\n## 安装\n\n```markdown\n## 代码中的标题\n```\n\n结尾正文", encoding="utf-8")
    chunks = split_documents(load_documents((tmp_path,)))
    assert len(chunks) == 1
    assert "## 代码中的标题" in chunks[0].page_content
    assert "结尾正文" in chunks[0].page_content


def test_pdf_without_text_requires_ocr(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(tmp_path / "scan.pdf")
    with pytest.raises(ValueError, match="OCR"):
        load_documents((tmp_path,))


def test_article_reference_formatted_as_heading_stays_in_original_clause(tmp_path):
    (tmp_path / "law.md").write_text(
        "# 法规\n\n#### 第九条\n\n除《电子商务法》\n\n#### 第十条\n\n规定的情形外。"
        "\n\n#### 第十条\n\n登记要求。", encoding="utf-8")
    chunks = split_documents(load_documents((tmp_path,)))
    assert len(chunks) == 2
    assert "规定的情形外。" in chunks[0].page_content
    assert "登记要求。" in chunks[1].page_content
